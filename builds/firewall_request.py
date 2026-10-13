"""Builds the "New Build Firewall Request Form" spreadsheet used for a standard
build firewall ticket, filling in the IP/HOSTNAME/REASON marker cells of the
site-specific xlsx template under builds/firewall_templates/. Pure spreadsheet
logic only — no Jira calls (see builds/jira_client.py / builds/firewall_ticket.py).
Ported from CASBAH-PrepareNewBuild's firewall_request.py.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

TEMPLATES_DIR = Path(__file__).parent / "firewall_templates"

SITES = {
    "rcd09": {
        "label": "RCD",
        "template": TEMPLATES_DIR / "RCD_New_Build_Firewall Request Form.xlsx",
        "ip_prefix": "10.192.",
    },
    "sjc17": {
        "label": "SJC",
        "template": TEMPLATES_DIR / "SJC_New_Build_Firewall Request Form.xlsx",
        "ip_prefix": "10.196.",
    },
}


class FirewallRequestError(Exception):
    """Raised when a firewall request spreadsheet can't be built for a site."""


def site_config(site: str) -> dict:
    """Look up the SITES entry for a site code (e.g. 'rcd09'), raising
    FirewallRequestError if there's no template configured for it."""
    config = SITES.get(site)
    if config is None:
        supported = ", ".join(SITES)
        raise FirewallRequestError(
            f"No firewall request template configured for site '{site}' "
            f"(supported: {supported})."
        )
    return config


def build_workbook(hostname: str, ip: str, reason: str, site: dict):
    """Fill in the IP/HOSTNAME/REASON marker cells of the site's template workbook.
    Formatting from the template is left untouched."""
    wb = load_workbook(site["template"])
    ws = wb.active
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == "IP":
                cell.value = ip
            elif cell.value == "HOSTNAME":
                cell.value = hostname
            elif cell.value == "REASON":
                cell.value = reason
    return wb


def build_filename(hostname: str, site: dict) -> str:
    return f"{site['label']}_New_Build_{hostname}_Firewall Request Form.xlsx"
