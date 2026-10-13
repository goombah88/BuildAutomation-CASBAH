"""One-off script generating the CASBAH demo talking-points deck. Not part of the
Django app; run manually with `python scripts/build_demo_deck.py` whenever the
talking points need regenerating. Requires python-pptx (dev-only, not in
requirements.txt since the running app never imports it)."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

TITLE_COLOR = RGBColor(0x1F, 0x2A, 0x44)
ACCENT_COLOR = RGBColor(0x2F, 0x6F, 0xED)

SLIDES = [
    (
        "Project CASBAH",
        [
            "CLL's Assisted Server Build Automated Helper",
            "One form submission replaces manually bouncing between "
            "GitLab, NetBox, vCenter, and hand-running terraform",
        ],
    ),
    (
        "The problem it solves",
        [
            "Server builds used to require manually touching 3-4 separate tools",
            "Manual steps are where typos, forgotten CI variables, and "
            "wrong permissions creep in",
            "CASBAH automates the same steps identically every time",
        ],
    ),
    (
        "Architecture & tech stack",
        [
            "Django 6.1, split into two independent long-running processes: "
            "web server + background worker",
            "django-q2 background task queue \u2014 uses the database itself "
            "as the queue, no Redis to stand up",
            "HTMX + Alpine.js via CDN \u2014 no Node build step, small footprint",
            "Dedicated clients for GitLab, NetBox, vCenter, and HashiCorp Vault",
        ],
    ),
    (
        "End-to-end pipeline automation",
        [
            "Form submit \u2192 background task creates the GitLab repo "
            "(real git clone/push)",
            "\u2192 creates the NetBox VM record",
            "\u2192 GitLab CI runs terraform plan",
            "\u2192 CASBAH auto-merges to main once plan is clean",
            "\u2192 terraform apply auto-triggers once the post-merge plan is clean",
            "What used to need 3-4 manual approval clicks is now hands-off",
        ],
    ),
    (
        "Security",
        [
            "Zero secrets in code or config \u2014 every credential comes from "
            "Vault at runtime via AppRole auth",
            "Per-build least-privilege tokens: each pipeline gets a short-lived "
            "(20-minute) Vault child token",
            "Real AD/LDAP authentication, gated to a specific AD group",
            "Fails closed: any misconfiguration or LDAP error denies login, "
            "never defaults to letting someone in",
            "Whole app requires login \u2014 no anonymous access",
            "Credential redaction: tokens are scrubbed before they can be "
            "logged, stored, or emailed",
        ],
    ),
    (
        "Reliability",
        [
            "233 automated tests, fully offline/mocked, run in under a second",
            "Also validated against real infrastructure end-to-end before "
            "being called \u201cdone\u201d",
            "No silent failures \u2014 partial failures show exactly which "
            "steps completed, nothing auto-rolled-back or hidden",
            "Full audit history on every build request, tied to real AD identity",
        ],
    ),
    (
        "Honest gaps / what's next",
        [
            "Currently on manage.py runserver / SQLite for dev",
            "Production deployment (systemd services or k8s, Postgres, "
            "TLS/reverse proxy) is the next phase",
            "Decommission / VM lifecycle management is explicitly out of "
            "scope for v1",
        ],
    ),
]


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for title, bullets in SLIDES:
        slide = prs.slides.add_slide(blank_layout)

        title_box = slide.shapes.add_textbox(
            Inches(0.6), Inches(0.4), Inches(12.1), Inches(1.0)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_run = title_frame.paragraphs[0].runs[0]
        title_run.font.size = Pt(36)
        title_run.font.bold = True
        title_run.font.color.rgb = TITLE_COLOR

        body_box = slide.shapes.add_textbox(
            Inches(0.9), Inches(1.6), Inches(11.5), Inches(5.5)
        )
        body_frame = body_box.text_frame
        body_frame.word_wrap = True
        for i, bullet in enumerate(bullets):
            paragraph = (
                body_frame.paragraphs[0] if i == 0 else body_frame.add_paragraph()
            )
            paragraph.text = f"\u2022  {bullet}"
            paragraph.font.size = Pt(20)
            paragraph.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            paragraph.space_after = Pt(14)

        accent_bar = slide.shapes.add_shape(
            1, Inches(0.6), Inches(1.35), Inches(2.0), Pt(4)
        )
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = ACCENT_COLOR
        accent_bar.line.fill.background()

    prs.save("CASBAH_Demo_Talking_Points.pptx")
    print("Wrote CASBAH_Demo_Talking_Points.pptx")


if __name__ == "__main__":
    build()
